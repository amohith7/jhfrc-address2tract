"""Tests for de-identified sharing and uniform column formatting."""

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.utils.io import (
    normalize_text_columns,
    share_file_path,
    write_output,
    write_share_file,
)


def test_normalize_strips_whitespace_and_excel_dot_zero():
    df = pd.DataFrame(
        {"id": ["  001  ", "30720.0", np.nan], "zip": ["37402.0", " 37403 ", "x.0"]}
    )
    out = normalize_text_columns(df)
    assert out["id"].tolist()[:2] == ["001", "30720"]
    assert pd.isna(out["id"].iloc[2])  # blank stays NaN (reject logic intact)
    assert out["zip"].tolist() == ["37402", "37403", "x.0"]  # only int.0 stripped


def test_share_file_drops_address_and_pii_columns(tmp_path):
    df = pd.DataFrame(
        {
            "client_id": ["001", "002"],
            "full_address": ["123 Main St, Chattanooga, 37402", "9 Oak Ave"],
            "cleaned_address": ["123 MAIN ST", "9 OAK AVE"],
            "census_tract_geoid": ["47065003100", "47065010800"],
            "match_status": ["Matched", "Matched_External"],
            "error_reason": [None, None],
        }
    )
    out = tmp_path / "results.xlsx"
    share = write_share_file(df, str(out), "client_id")

    assert share == share_file_path(str(out))
    got = pd.read_excel(share, dtype=str)
    assert list(got.columns) == ["client_id", "census_tract_geoid", "match_status"]
    assert "full_address" not in got.columns
    assert "cleaned_address" not in got.columns
    assert got["client_id"].tolist() == ["001", "002"]  # leading zero preserved


def test_key_columns_written_as_text_for_xlookup(tmp_path):
    df = pd.DataFrame({"client_id": ["001"], "census_tract_geoid": ["47065003100"]})
    out = tmp_path / "keys.xlsx"
    write_output(df, str(out), text_columns=["client_id", "census_tract_geoid"])
    ws = load_workbook(out).active
    # data row (row 2) of both key columns must be Text format ("@")
    assert ws.cell(row=2, column=1).number_format == "@"
    assert ws.cell(row=2, column=2).number_format == "@"


def test_share_file_works_for_csv(tmp_path):
    df = pd.DataFrame(
        {
            "id": ["001"],
            "full_address": ["secret 123 Main St"],
            "census_tract_geoid": ["47065003100"],
            "match_status": ["Matched"],
        }
    )
    out = tmp_path / "results.csv"
    share = write_share_file(df, str(out), "id")
    text = share.read_text()
    assert "Main St" not in text  # no address leaked
    assert "47065003100" in text
